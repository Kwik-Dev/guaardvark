"""Build a fine-tuning dataset from a folder of a person's own notes.

The stages, in order, each reported with counts:

1. scan: the folder is walked with Bulk Import's rules (scan_import_source:
   no hidden folders or files, no credential files, nothing of Guaardvark's
   own data), keeping the chosen file kinds;
2. scrub: always on (folder_scrub), before any text reaches a model;
3. chunk: EnhancedRAGChunker, strategy auto, as indexing does;
4. pairs: questions and answers from the local model (pair_generator);
5. filter: duplicate questions, answers the chunk does not support, the
   heuristic quality rules and length bounds;
6. write: one .jsonl in the training datasets folder and its build report.

The check (scan and scrub, no model) and a three-chunk timing probe give the
person the scrub report and an estimate before anything is built.
"""
from __future__ import annotations

import hashlib
import json
import logging
import math
import os
import random
import re
import time
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Callable, Iterable, Optional

from backend.services.training.folder_scrub import Scrubber, load_terms
from backend.services.training.pair_generator import (
    QUESTION_KINDS,
    PairGenerator,
    as_training_row,
)

logger = logging.getLogger(__name__)

DEFAULT_KINDS = (".md", ".txt")
ALL_KINDS = (".md", ".txt", ".html", ".docx", ".pdf")
DEFAULT_PAIRS_PER_1000_WORDS = 3
MAX_PAIRS_PER_1000_WORDS = 20
PROBE_CHUNKS = 3
# A file this long is a tool dump or a word list (a grep output, a tokenizer
# vocabulary), not notes: checked on a real notes folder on 2026-10-12, its
# seven files over this held 30 % of the words and none was prose notes,
# while the notes themselves sat well under it. Left out unless ticked.
LARGE_FILE_WORDS = 50_000
# Passages are chosen across files in this seeded order, so a rebuild picks the same ones.
SELECTION_SEED = 42
# A file gives at most this many times its fair share of the chosen passages
# (and never fewer than PER_FILE_MIN), so one long file cannot fill the dataset.
PER_FILE_SHARE = 3
PER_FILE_MIN = 4
# Pairs lost to the filter, as headroom on top of what training can use.
FILTER_HEADROOM = 1.25
# A chunk's size before the probe has measured one; the chunker's own target.
ASSUMED_WORDS_PER_CHUNK = 250
# A passage shorter than this is joined to the next one from the same file. The
# indexing chunker splits notes at every heading, so on a real notes folder on
# 2026-10-12 18 % of its 31,141 chunks were under 50 words and 37 % under 100:
# fragments too short to ask a question that needs the passage to answer.
MIN_PASSAGE_WORDS = 120

# Filter bounds. Answers must share this much of their content words with the
# text they came from (two passages for multi-hop, a summary paraphrases more).
SUPPORT_MIN = {"specific": 0.6, "reasoning": 0.5, "multi_hop": 0.4, "summary": 0.4}
QUESTION_CHARS = (10, 400)
ANSWER_CHARS = (20, 2000)
QUALITY_MIN = 0.5


class BuildCancelled(Exception):
    """The person cancelled the build; nothing is written."""


def training_dir() -> Path:
    from backend.config import STORAGE_DIR

    return Path(STORAGE_DIR) / "training"


def scrub_terms_path() -> Path:
    return training_dir() / "scrub_terms.txt"


def build_reports_dir() -> Path:
    # Not beside the dataset: the dataset readers take every .json in the
    # datasets folder as rows, so a report there would read as a broken dataset.
    path = training_dir() / "build_reports"
    path.mkdir(parents=True, exist_ok=True)
    return path


def normalise_kinds(kinds) -> list[str]:
    chosen = []
    for kind in kinds or DEFAULT_KINDS:
        kind = str(kind).lower().strip()
        kind = kind if kind.startswith(".") else f".{kind}"
        if kind in ALL_KINDS and kind not in chosen:
            chosen.append(kind)
    return chosen or list(DEFAULT_KINDS)


# Training-only rules on top of Bulk Import's: a folder of notes often holds
# copied software too (a dependency tree, a virtual environment, a copied app
# with its tokenizer files), whose .md and .txt files are not anyone's notes.
# Folders pruned by name at any depth:
PACKAGE_DIRS = {
    "node_modules", "site-packages", "dist-packages", "vendor", "third_party", "third-party",
    "bower_components", "build", "dist", "__pycache__", "venv", "virtualenv",
}
PACKAGE_DIR_SUFFIXES = (".egg-info", ".dist-info")
# Files that sit among notes but are not notes: tokenizer vocabularies,
# dependency pins, licences and package metadata.
# Names only: a note called "license notes for the shop" stays.
NOT_NOTES_NAME = re.compile(
    r"^(?:"
    r"(?i:merges|vocab|requirements|constraints)(?:[._-][\w.-]*)?\.(?i:txt)"
    r"|(?i:licen[cs]e|copying|notice|authors)(?:[._-][A-Z0-9.]{1,15})?\.(?i:txt|md)"
    r"|(?i:cmakelists|dependency_links|top_level|sources|entry_points|installed-files)\.(?i:txt)"
    r")$")
# Files whose names say they hold a credential. Left out whole and listed by
# name, since their content cannot be trusted to match any scrub rule.
CREDENTIAL_NAME = re.compile(
    r"\b(?:api ?keys?|apikeys?|recovery (?:codes?|keys?)|passwords?|passwd|secrets?|tokens?"
    r"|credentials?|bitlocker|private ?keys?)\b", re.I)


def _package_dir(rel: str) -> Optional[str]:
    """The topmost copied-software folder a file sits in, or None."""
    parts = Path(rel).parts[:-1]
    for i, part in enumerate(parts):
        low = part.lower()
        if low in PACKAGE_DIRS or low.endswith(PACKAGE_DIR_SUFFIXES):
            return str(Path(*parts[:i + 1]))
    return None


def is_credential_name(name: str) -> bool:
    return bool(CREDENTIAL_NAME.search(re.sub(r"[_\-.]+", " ", Path(name).stem)))


def not_prose_reason(text: str) -> Optional[str]:
    """Why a file's text is not prose to learn from (a token list, minified or
    encoded text), or None. Cheap, on the first 50k characters."""
    sample = text[:50000]
    if not sample.strip():
        return None
    letters_spaces = sum(1 for c in sample if c.isalpha() or c.isspace())
    if letters_spaces / len(sample) < 0.6:
        return "mostly symbols or digits"
    if any(len(line) > 1000 and line.count(" ") < len(line) * 0.02 for line in sample.splitlines()):
        return "long unbroken lines"
    lines = [line for line in sample.splitlines() if line.strip()]
    if len(lines) >= 200:
        short = sum(1 for line in lines if len(line.split()) <= 2)
        sentences = len(re.findall(r"[.!?](?:\s|$)", sample))
        if short / len(lines) > 0.9 and sentences < len(sample.split()) / 50:
            return "a list of single words or tokens"
    return None


def scan(folder: str, kinds) -> tuple[list[tuple[str, str]], dict]:
    """(files, skipped): (full path, path relative to the folder) for each file
    of the chosen kinds that Bulk Import's rules and the training rules above
    let in. skipped["credential_files"] names the files left out for their name."""
    from backend.celery_tasks_isolated import _guaardvark_storage_dirs, scan_import_source

    kinds = normalise_kinds(kinds)
    storage_dir, upload_dir, backup_dir = _guaardvark_storage_dirs()
    found, skipped = scan_import_source(folder, set(kinds), storage_dir, upload_dir,
                                        backup_dir=backup_dir)
    skipped.update(other_kinds=0, package_files=0, not_notes=0, credential_names=0)
    package_dirs, credential_files, files = set(), [], []
    for full, rel, name in found:
        if os.path.splitext(name)[1].lower() not in kinds:
            skipped["other_kinds"] += 1
        elif (package := _package_dir(rel)) is not None:
            package_dirs.add(package)
            skipped["package_files"] += 1
        elif NOT_NOTES_NAME.match(name):
            skipped["not_notes"] += 1
        elif is_credential_name(name):
            skipped["credential_names"] += 1
            credential_files.append(rel)
        else:
            files.append((full, rel))
    skipped["package_dirs"] = len(package_dirs)
    skipped["credential_files"] = sorted(credential_files)
    return files, skipped


def read_text(path: str) -> str:
    suffix = os.path.splitext(path)[1].lower()
    if suffix == ".docx":
        from docx import Document

        return "\n".join(p.text for p in Document(path).paragraphs)
    if suffix == ".pdf":
        # The plain-text reader the indexer uses for PDFs (indexing_service's
        # PDFReaderClass), not Docling: text only, and fast enough for a check.
        from llama_index.readers.file import PDFReader

        return "\n".join(doc.text for doc in PDFReader().load_data(file=Path(path)))
    with open(path, encoding="utf-8", errors="replace") as f:
        text = f.read()
    if suffix == ".html":
        from bs4 import BeautifulSoup

        text = BeautifulSoup(text, "html.parser").get_text("\n")
    return text


def _words(text: str) -> int:
    return len(text.split())


def prepare(folder: str, kinds, terms: Optional[list[str]] = None,
            include_large: Iterable[str] = ()) -> dict:
    """Scan, read and scrub. Returns the scrubbed texts and the report parts.
    Files over LARGE_FILE_WORDS are listed in skipped["large_files"] and left
    out unless their path is in include_large."""
    files, skipped = scan(folder, kinds)
    include_large = set(include_large or ())
    skipped["not_prose"] = 0
    skipped["large_files"] = []
    terms = terms if terms is not None else load_terms(scrub_terms_path())
    scrubber = Scrubber(terms)
    # File names carry the same names as the text (a client's name in a memo's
    # name), and a summary question quotes the file name, so the path a doc
    # carries into chunks, questions and row metadata is scrubbed too. Its own
    # scrubber keeps the report's counts about the text.
    path_scrubber = Scrubber(terms)
    docs, unreadable = [], []
    for full, rel in files:
        try:
            text = read_text(full)
        except Exception as e:
            unreadable.append(f"{rel}: {e}")
            continue
        if not text.strip():
            if rel.lower().endswith(".pdf"):
                unreadable.append(f"{rel}: no text layer")
            continue
        if not_prose_reason(text):
            skipped["not_prose"] += 1
            continue
        words = _words(text)
        if words > LARGE_FILE_WORDS and rel not in include_large:
            skipped["large_files"].append({"source_path": rel, "words": words})
            continue
        clean = scrubber.scrub(text, rel)
        if clean.strip():
            docs.append({"source_path": path_scrubber.scrub(rel), "text": clean, "words": _words(clean)})
    return {
        "folder": folder,
        "kinds": normalise_kinds(kinds),
        "docs": docs,
        "files": len(docs),
        "words": sum(d["words"] for d in docs),
        "skipped": skipped,
        "unreadable": unreadable,
        "scrub": scrubber.report(),
    }


def _join_short(texts: list[str]) -> list[str]:
    """Join each passage under MIN_PASSAGE_WORDS to the ones after it in the same
    file until it reaches that size; a short tail joins the passage before it."""
    joined: list[str] = []
    current = ""
    for text in texts:
        current = f"{current}\n\n{text}" if current else text
        if _words(current) >= MIN_PASSAGE_WORDS:
            joined.append(current)
            current = ""
    if current:
        if joined:
            joined[-1] = f"{joined[-1]}\n\n{current}"
        else:
            joined.append(current)
    return joined


def chunk_docs(docs: list[dict], limit: Optional[int] = None) -> list[dict]:
    """Chunks as indexing makes them, short ones joined (_join_short), each with
    its file and hash."""
    from llama_index.core import Document as LlamaDocument
    from backend.utils.enhanced_rag_chunking import EnhancedRAGChunker

    chunker = EnhancedRAGChunker()
    chunks = []
    for doc in docs:
        nodes = chunker.chunk_documents([LlamaDocument(text=doc["text"], metadata={})], strategy_name="auto")
        texts = [n.get_content() for n in (nodes or []) if getattr(n, "text", None)] or [doc["text"][:2000]]
        texts = _join_short(texts)
        for i, text in enumerate(texts):
            chunks.append({
                "source_path": doc["source_path"],
                "text": text,
                "words": _words(text),
                "hash": hashlib.sha256(text.encode()).hexdigest(),
                # The passage after it in the same file, for a multi-hop question.
                "next_text": texts[i + 1] if i + 1 < len(texts) else None,
            })
            if limit and len(chunks) >= limit:
                return chunks
    return chunks


def questions_for(chunk_words: int, per_1000: float) -> int:
    """Questions asked of one chunk; its summary pair comes on top."""
    return max(1, round(chunk_words * per_1000 / 1000))


def pairs_per_passage(chunk_words: int, per_1000: float) -> int:
    return questions_for(chunk_words, per_1000) + 1


# One pass over ~8,000 varied rows covers more of a folder and overfits less
# than three passes over the 2,700 that would fill the step cap three times.
def enough_to_train() -> int:
    """Pairs a default fine-tune can use, plus headroom for the filter.

    The trainer stops at DEFAULT_STEPS["max"] steps of DEFAULT_BATCH_SIZE x
    GRADIENT_ACCUMULATION_STEPS rows, and holds out at most EVAL_SPLIT["max_rows"]
    rows; a row beyond that is never seen. Rounded up to a multiple of 500."""
    from backend.tasks.training_tasks import (
        DEFAULT_BATCH_SIZE, DEFAULT_STEPS, EVAL_SPLIT, GRADIENT_ACCUMULATION_STEPS)

    rows = DEFAULT_STEPS["max"] * DEFAULT_BATCH_SIZE * GRADIENT_ACCUMULATION_STEPS + EVAL_SPLIT["max_rows"]
    return math.ceil(rows * FILTER_HEADROOM / 500) * 500


TARGET_CHOICES = ("enough", 5000, 2000, "all")


def resolve_target(target) -> Optional[int]:
    """Pairs to aim for: a number, "enough" (the default) or "all" (None)."""
    if target in (None, "enough"):
        return enough_to_train()
    if target == "all":
        return None
    return int(target)


def _selection_counts(available: dict[str, int], needed: int) -> tuple[list[str], dict[str, int], int]:
    """How many passages each file gives toward `needed`: files in a seeded
    order, one passage from each in turn, so every file gives one before any
    gives two, and none above its cap (PER_FILE_SHARE times its fair share,
    at least PER_FILE_MIN). Returns (file order, counts, cap)."""
    files = sorted(available)
    random.Random(SELECTION_SEED).shuffle(files)
    cap = max(PER_FILE_MIN, math.ceil(PER_FILE_SHARE * needed / max(1, len(files))))
    taken, total = Counter(), 0
    while total < needed:
        progressed = False
        for f in files:
            if total >= needed:
                break
            if taken[f] < min(cap, available[f]):
                taken[f] += 1
                total += 1
                progressed = True
        if not progressed:
            break
    return files, taken, cap


def select_passages(chunks: list[dict], target_pairs: Optional[int], per_1000: float) -> tuple[list[dict], dict]:
    """The passages to write pairs for, spread across files (_selection_counts);
    within a file they come in a seeded order, so they spread over it. When the
    per-file cap stops it short of the target, the report says by how much.
    No target: every passage."""
    by_file: dict[str, list[dict]] = {}
    for c in chunks:
        by_file.setdefault(c["source_path"], []).append(c)
    info = {"target_pairs": target_pairs, "passages_available": len(chunks), "files": len(by_file)}
    if not chunks or target_pairs is None:
        return chunks, {**info, "passages_chosen": len(chunks), "files_contributing": len(by_file)}
    expected = sum(pairs_per_passage(c["words"], per_1000) for c in chunks) / len(chunks)
    needed = min(len(chunks), math.ceil(target_pairs / expected))
    files, taken, cap = _selection_counts({f: len(v) for f, v in by_file.items()}, needed)
    rng = random.Random(SELECTION_SEED)
    chosen = []
    for f in files:
        queue = list(by_file[f])
        rng.shuffle(queue)
        chosen.extend(queue[:taken[f]])
    position = {id(c): i for i, c in enumerate(chunks)}
    chosen.sort(key=lambda c: position[id(c)])
    return chosen, {**info, "passages_needed": needed, "passages_chosen": len(chosen),
                    "files_contributing": sum(1 for f in taken if taken[f]), "per_file_cap": cap,
                    "short_by_cap": needed - len(chosen)}


def estimate(prepared: dict, per_1000: float, probe: Optional[dict] = None, target="enough") -> dict:
    """Passages, pairs and time before a build, for the chosen target and for
    each choice the dialog offers. Without a probe the passage size is assumed
    and there is no time; the probe measures both."""
    measured = bool(probe and probe.get("words_per_chunk"))
    words_per_chunk = probe["words_per_chunk"] if measured else ASSUMED_WORDS_PER_CHUNK
    if measured and probe.get("passages_per_file"):
        per_file = dict(probe["passages_per_file"])
        per_passage = probe.get("pairs_per_passage") or pairs_per_passage(round(words_per_chunk), per_1000)
    else:
        per_file = {d["source_path"]: max(1, math.ceil(d["words"] / words_per_chunk)) for d in prepared["docs"]}
        per_passage = pairs_per_passage(round(words_per_chunk), per_1000)
    available = sum(per_file.values())
    seconds_per = probe.get("seconds_per_chunk") if probe else None

    def plan(target_pairs):
        if target_pairs is None or not per_file:
            passages = available
        else:
            _, taken, _ = _selection_counts(per_file, min(available, math.ceil(target_pairs / per_passage)))
            passages = sum(taken.values())
        out = {"target_pairs": target_pairs, "passages": passages,
               "pairs_before_filter": round(passages * per_passage)}
        if seconds_per:
            out["seconds"] = round(passages * seconds_per)
        return out

    chosen = plan(resolve_target(target))
    return {
        "pairs_per_1000_words": per_1000,
        "words_per_chunk": words_per_chunk,
        "basis": "measured by the probe" if measured else f"assumed {ASSUMED_WORDS_PER_CHUNK} words per chunk",
        "passages_available": available,
        "chunks": chosen["passages"],
        "pairs_before_filter": chosen["pairs_before_filter"],
        **({"seconds": chosen["seconds"]} if "seconds" in chosen else {}),
        "target_pairs": chosen["target_pairs"],
        "enough_to_train": enough_to_train(),
        "choices": {str(c): plan(resolve_target(c)) for c in TARGET_CHOICES},
        "estimated": True,
    }


def probe_timing(prepared: dict, per_1000: float, generator: Optional[PairGenerator] = None,
                 target="enough") -> dict:
    """Generate (and throw away) the pairs of a few passages to measure the time.

    The whole folder is chunked, so the passage counts are real rather than
    projected from a few passages, and the timed passages are spread over the
    ones a build of `target` would choose: the first passages of one file are
    not a fair sample of a folder (one transcript's short turns timed 48 words
    a passage on a folder whose passages average 212)."""
    all_chunks = chunk_docs(prepared["docs"])
    if not all_chunks:
        return {"chunks": 0, "reason": "nothing to read in the folder"}
    chosen, _ = select_passages(all_chunks, resolve_target(target), per_1000)
    pool = chosen or all_chunks
    step = max(1, len(pool) // PROBE_CHUNKS)
    sample = pool[::step][:PROBE_CHUNKS]
    generator = generator or PairGenerator()
    started = time.monotonic()
    pairs = 0
    for chunk in sample:
        pairs += len(_pairs_for_chunk(generator, chunk, per_1000, 0))
    seconds = time.monotonic() - started
    model = getattr(generator._llm, "model", None)
    per_file = Counter(c["source_path"] for c in all_chunks)
    return {
        "chunks": len(sample),
        "pairs": pairs,
        "seconds": round(seconds, 1),
        "seconds_per_chunk": round(seconds / len(sample), 1),
        "words_per_chunk": round(sum(c["words"] for c in all_chunks) / len(all_chunks), 1),
        "pairs_per_passage": round(sum(pairs_per_passage(c["words"], per_1000) for c in all_chunks)
                                   / len(all_chunks), 2),
        "passages_per_file": dict(per_file),
        "model": model,
    }


def _pairs_for_chunk(generator: PairGenerator, chunk: dict, per_1000: float, kind_offset: int) -> list[dict]:
    pairs = []
    for i in range(questions_for(chunk["words"], per_1000)):
        kind = QUESTION_KINDS[(kind_offset + i) % len(QUESTION_KINDS)]
        other = chunk.get("next_text") if kind == "multi_hop" else None
        pair = generator.question(chunk["text"], kind, other)
        if pair:
            pair["support_text"] = chunk["text"] + ("\n" + other if other else "")
            pairs.append(pair)
    summary = generator.summary(chunk["text"], Path(chunk["source_path"]).stem)
    if summary:
        summary["support_text"] = chunk["text"]
        pairs.append(summary)
    for pair in pairs:
        pair["source_path"], pair["chunk_hash"] = chunk["source_path"], chunk["hash"]
    return pairs


def generate(chunks: list[dict], per_1000: float, generator: PairGenerator,
             should_stop: Callable[[], bool] = lambda: False,
             on_progress: Callable[[int, int, int], None] = lambda done, total, pairs: None) -> list[dict]:
    """Pairs for every chunk. Stops between chunks when asked."""
    pairs = []
    for i, chunk in enumerate(chunks):
        if should_stop():
            raise BuildCancelled()
        pairs.extend(_pairs_for_chunk(generator, chunk, per_1000, i))
        on_progress(i + 1, len(chunks), len(pairs))
    return pairs


_CONTENT_WORD = re.compile(r"[a-z0-9][a-z0-9'\-]{3,}")


def _support(answer: str, text: str) -> float:
    words = set(_CONTENT_WORD.findall(answer.lower()))
    if not words:
        return 0.0
    source = set(_CONTENT_WORD.findall(text.lower()))
    return len(words & source) / len(words)


def _normalised(question: str) -> str:
    return re.sub(r"[^a-z0-9 ]", "", re.sub(r"\s+", " ", question.lower())).strip()


def filter_pairs(pairs: list[dict]) -> tuple[list[dict], dict]:
    """Keep pairs that are unique, supported by their text and of sound length."""
    from backend.tasks.training_tasks import PARSER_DIR
    import sys

    if str(PARSER_DIR) not in sys.path:
        sys.path.insert(0, str(PARSER_DIR))
    from quality_filter import QualityFilter

    quality = QualityFilter()
    dropped = {"duplicate": 0, "unsupported": 0, "length": 0, "quality": 0}
    seen, kept = set(), []
    for pair in pairs:
        q, a = pair["question"], pair["answer"]
        key = _normalised(q)
        if key in seen:
            dropped["duplicate"] += 1
            continue
        if not (QUESTION_CHARS[0] <= len(q) <= QUESTION_CHARS[1]
                and ANSWER_CHARS[0] <= len(a) <= ANSWER_CHARS[1]):
            dropped["length"] += 1
            continue
        if _support(a, pair["support_text"]) < SUPPORT_MIN.get(pair["kind"], 0.5):
            dropped["unsupported"] += 1
            continue
        score, _issues = quality.score_example(q, a)
        if score < QUALITY_MIN:
            dropped["quality"] += 1
            continue
        seen.add(key)
        kept.append(pair)
    return kept, dropped


def dataset_name(folder: str) -> str:
    stem = re.sub(r"[^A-Za-z0-9_-]+", "_", Path(folder.rstrip("/")).name).strip("_") or "folder"
    return f"folder_{stem}_{datetime.now().strftime('%Y%m%d_%H%M%S')}"


def write(pairs: list[dict], report: dict, name: str) -> tuple[str, str]:
    """The dataset in the training datasets folder, the report in build_reports."""
    from backend.tasks.training_tasks import parsed_datasets_dir

    out_path = parsed_datasets_dir() / f"{name}.jsonl"
    with open(out_path, "w", encoding="utf-8") as f:
        for pair in pairs:
            f.write(json.dumps(as_training_row(pair, pair["source_path"], pair["chunk_hash"])) + "\n")
    report_path = build_reports_dir() / f"{name}.json"
    report_path.write_text(json.dumps(report, indent=2), encoding="utf-8")
    return str(out_path), str(report_path)
